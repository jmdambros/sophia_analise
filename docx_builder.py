"""
docx_builder.py
---------------
Converte o rascunho de peça processual (texto simples devolvido pela IA,
com a convenção leve de formatação descrita em DRAFT_SUFFIX de
analysis.py) num arquivo .docx já formatado, para que o advogado não
precise copiar e colar manualmente.
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt


def build_docx_from_draft(tipo_peca: str, texto: str) -> bytes:
    """Gera os bytes de um .docx a partir do texto plano retornado pela IA.

    Convenção esperada no texto:
      - Linha começando com "# "  -> título principal (Heading 1)
      - Linha começando com "## " -> título de seção (Heading 2)
      - Linha em branco           -> ignorada (só separa parágrafos)
      - Qualquer outra linha      -> parágrafo normal
    """
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    aviso = doc.add_paragraph()
    aviso_run = aviso.add_run(
        "MINUTA GERADA POR IA — REVISAR INTEGRALMENTE ANTES DE PROTOCOLAR."
    )
    aviso_run.bold = True
    aviso_run.font.size = Pt(10)
    doc.add_paragraph()

    tem_titulo = False
    for raw_line in texto.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            tem_titulo = True
        else:
            doc.add_paragraph(line)

    if not tem_titulo:
        # Defesa: se o modelo não seguiu a convenção de "# título", garante
        # pelo menos um cabeçalho com o tipo de peça pedido.
        doc.add_heading((tipo_peca or "MINUTA").upper(), level=1)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()